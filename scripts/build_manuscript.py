"""Build editable Word equations with Pandoc, then format the Russian manuscript.

Requires pandoc and python-docx. Render and inspect the output before submission.
"""
import argparse
from pathlib import Path
import subprocess
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",required=True)
    args=parser.parse_args()
    source=Path(__file__).resolve().parents[1]/"paper"/"article_ru.md"
    out=Path(args.output).resolve()
    out.parent.mkdir(parents=True,exist_ok=True)
    subprocess.run(["pandoc",str(source),"--from=markdown+tex_math_dollars","--standalone",
        "--resource-path",str(source.parent),"-o",str(out)],check=True)
    doc=Document(out)
    # Literal square brackets avoid a LibreOffice import error in which a
    # stretchy closing bracket is rendered as a parenthesis. Keep native math.
    for delimiter in reversed(doc.element.xpath('.//m:d')):
        props=delimiter.find(qn("m:dPr"))
        if props is None: continue
        begin=props.find(qn("m:begChr")); end=props.find(qn("m:endChr"))
        if begin is None or end is None: continue
        if (begin.get(qn("m:val")),end.get(qn("m:val"))) != ("[","]"): continue
        entries=delimiter.findall(qn("m:e"))
        if len(entries)!=1: continue
        brackets=[]
        for value in ("[","]"):
            run=OxmlElement("m:r")
            props_run=OxmlElement("m:rPr")
            style=OxmlElement("m:sty"); style.set(qn("m:val"),"p")
            props_run.append(style); run.append(props_run)
            text=OxmlElement("m:t"); text.text=value; run.append(text)
            brackets.append(run)
        parent=delimiter.getparent(); position=parent.index(delimiter)
        children=[brackets[0],*list(entries[0]),brackets[1]]
        parent.remove(delimiter)
        for offset,child in enumerate(children): parent.insert(position+offset,child)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(0.65); sec.bottom_margin=Inches(0.65)
    sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8)
    for s in doc.styles:
        if s.type==1:
            s.font.name="Times New Roman"
            fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in list(fonts.attrib):
                if attr.endswith("Theme"): del fonts.attrib[attr]
            for attr in ("ascii","hAnsi","eastAsia","cs"):
                fonts.set(qn(f"w:{attr}"),"Times New Roman")
            s.font.size=Pt(11.5)
            s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.line_spacing=1.05
            s.paragraph_format.space_after=Pt(5)
    for name in ("Normal","Body Text","First Paragraph"):
        s=doc.styles[name]
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        s.paragraph_format.first_line_indent=Inches(0.22)
        s.paragraph_format.widow_control=True
    title=doc.styles["Title"]
    title.font.size=Pt(16); title.font.bold=True
    title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after=Pt(6)
    title.paragraph_format.first_line_indent=Inches(0)
    for name in ("Author","Subtitle"):
        if name in doc.styles:
            doc.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            doc.styles[name].paragraph_format.first_line_indent=Inches(0)
    for i in range(1,4):
        s=next(style for style in doc.styles if style.style_id==f"Heading{i}")
        s.font.size=Pt(12); s.font.bold=True
        s.paragraph_format.space_before=Pt(10)
        s.paragraph_format.space_after=Pt(5)
        s.paragraph_format.keep_with_next=True
        s.paragraph_format.first_line_indent=Inches(0)
    paragraphs=doc.paragraphs
    for position,p in enumerate(paragraphs):
        if position+1 < len(paragraphs) and paragraphs[position+1]._p.xpath('.//m:oMathPara'):
            p.paragraph_format.keep_with_next=True
        if p.style.style_id.startswith("Heading") or p.style.name=="Title":
            for r in p.runs:
                r.font.name="Times New Roman"
                r.bold=True
                fonts=r._r.get_or_add_rPr().get_or_add_rFonts()
                for attr in list(fonts.attrib):
                    if attr.endswith("Theme"): del fonts.attrib[attr]
                fonts.set(qn("w:ascii"),"Times New Roman")
                fonts.set(qn("w:hAnsi"),"Times New Roman")
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.space_after=Pt(6)
        if p.text.startswith(("Аннотация.","Ключевые слова:")):
            p.paragraph_format.first_line_indent=Inches(0)
        if p.text.startswith("Таблица "):
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(7)
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(3)
        if "Caption" in p.style.name or p.text.startswith("Рисунок "):
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=False
            p.paragraph_format.space_after=Pt(7)
            for r in p.runs:r.font.size=Pt(10.5)
        if p.style.name=="Compact":
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.space_after=Pt(4)
            for r in p.runs: r.font.size=Pt(10.5)
    for table in doc.tables:
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        widths=([0.85,1.50,1.00,0.85,0.90,1.75] if len(table.columns)==6
                else [6.85/len(table.columns)]*len(table.columns))
        for i,w in enumerate(widths): table.columns[i].width=Inches(w)
        props=table._tbl.tblPr
        borders=OxmlElement("w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"D9D9D9"); borders.append(e)
        props.append(borders)
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            cant=OxmlElement("w:cantSplit"); trpr.append(cant)
            if ri==0: trpr.append(OxmlElement("w:tblHeader"))
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                for old in list(tcpr.findall(qn("w:tcBorders"))):tcpr.remove(old)
                cb=OxmlElement("w:tcBorders")
                for edge in ("top","left","bottom","right"):
                    e=OxmlElement(f"w:{edge}");e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");cb.append(e)
                tcpr.append(cb)
                margins=OxmlElement("w:tcMar")
                for edge in ("top","left","bottom","right"):
                    m=OxmlElement(f"w:{edge}"); m.set(qn("w:w"),"70"); m.set(qn("w:type"),"dxa"); margins.append(m)
                tcpr.append(margins)
                if ri==0:
                    shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"ECECEC"); tcpr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Inches(0)
                    p.paragraph_format.alignment=(WD_ALIGN_PARAGRAPH.CENTER if ri==0
                        else WD_ALIGN_PARAGRAPH.LEFT if ci<2 else WD_ALIGN_PARAGRAPH.RIGHT)
                    p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.line_spacing=1
                    for r in p.runs:
                        r.font.size=Pt(10.5)
                        if ri==0:r.bold=True
    # Remove inherited decorative borders from title and paragraph styles.
    for element in doc.styles.element.xpath('.//w:pBdr'):
        element.getparent().remove(element)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
    doc.core_properties.author="Салман Али"
    doc.core_properties.title="Влияние согласованности оптического потока на качество интерполяции видеокадров"
    doc.core_properties.subject="Результаты контролируемого эксперимента на 1240 тройках SNU-FILM"
    doc.save(out)
    print(out)


if __name__=="__main__":main()
